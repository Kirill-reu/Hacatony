"""ТЗ §2 item 3 core assembly: takes the user's uploaded template, a
``SlidePlan`` (content_generator) and the retrieved facts, and produces a
populated ``Presentation`` — new slides built on the template's *own*
layouts, with text placed via run-level edits (not ``text_frame.text = ...``,
which collapses inherited formatting), and charts/tables/pictures inserted
as native objects in place of the placeholder they replace, never as a
flattened image (required for the .pptx export to "count").

Three python-pptx facts drive the placeholder-matching approach here:
loading the template and clearing ``<p:sldIdLst>`` keeps every layout/master
(and their theme, fonts, colors) while dropping the template's own sample
slides; a fresh ``add_slide(layout)`` placeholder has no runs to preserve,
so the only formatting risk is on the title where we still edit the
existing run when the layout seeds one; and shapes leftover from a resolved
visual (a picture well with no image, an unused second content box) are
deleted outright rather than left empty, since an empty placeholder is
itself flagged by the audit.
"""
from __future__ import annotations

import logging
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.ns import qn

from ..llm.image_client import ImageClient
from ..llm.text_client import LLMClient
from .content_generator import SlidePlan, SlidePlanItem, VisualSpec
from .data_provider import SearchResult
from .template_parser import DesignSystem, LayoutSpec
from .visual_generator import (
    TableSpec,
    add_icon_row,
    build_chart_spec,
    build_table_spec,
    generate_image,
)

logger = logging.getLogger(__name__)

_CHART_TYPE_MAP = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}


# --------------------------------------------------------------------------
# Template slide stripping (keep masters/layouts/theme, drop sample slides)
# --------------------------------------------------------------------------


def _strip_existing_slides(prs: Presentation) -> None:
    xml_slides = prs.slides._sldIdLst
    for sld_id in list(xml_slides):
        r_id = sld_id.get(qn("r:id"))
        prs.part.drop_rel(r_id)
        xml_slides.remove(sld_id)


def _pick_accent(design_system: DesignSystem) -> str:
    for token in design_system.colors:
        if token.role.startswith("accent"):
            return token.hex
    return "4F81BD"


# --------------------------------------------------------------------------
# Placeholder matching (by idx, using the roles template_parser classified)
# --------------------------------------------------------------------------


_PICTURE_STENCIL_HINTS = ("вставить фото", "insert photo", "вставить изображение", "add photo", "add image")


def _find_picture_stencil_box(actual_layout) -> tuple[int, int, int, int] | None:
    """This template (like several decks authored in Google Slides) marks
    where a photo goes with an ordinary text box reading "Вставить фото" —
    not a real PICTURE-type placeholder. Because it isn't a placeholder,
    ``add_slide(layout)`` never copies it onto the new slide at all (only
    placeholders are; a layout's plain decorative/static shapes stay only
    on the layout and are otherwise just inherited visually) — so it's
    invisible to ``slide.shapes`` and, before this, silently ignored.
    Reads its position straight off the layout instead, so a real picture
    can be dropped in the same spot on the slide.
    """
    for shape in actual_layout.shapes:
        if shape.is_placeholder or not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip().lower()
        if any(hint in text for hint in _PICTURE_STENCIL_HINTS) and shape.width and shape.height:
            return shape.left or 0, shape.top or 0, shape.width, shape.height
    return None


def _match_slide_placeholders(slide, layout_spec: LayoutSpec):
    by_idx = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
    title = None
    bodies = []
    picture = None
    for spec in layout_spec.placeholders:
        ph = by_idx.get(spec.idx)
        if ph is None:
            continue
        if spec.type in ("TITLE", "CENTER_TITLE"):
            title = ph
        elif spec.type in ("BODY", "OBJECT", "SUBTITLE"):
            bodies.append(ph)
        elif spec.type == "PICTURE":
            picture = ph
    bodies.sort(key=lambda ph: ph.left or 0)
    return title, bodies, picture


def _remove_shape(shape) -> None:
    el = shape._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


# --------------------------------------------------------------------------
# Text placeholders
# --------------------------------------------------------------------------


def _set_title(placeholder, text: str) -> None:
    tf = placeholder.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    if p0.runs:
        run = p0.runs[0]
        for extra in list(p0.runs[1:]):
            extra._r.getparent().remove(extra._r)
    else:
        run = p0.add_run()
    run.text = text
    for extra_p in list(tf.paragraphs[1:]):
        extra_p._p.getparent().remove(extra_p._p)


def _set_bullets(placeholder, bullets: list[str]) -> None:
    tf = placeholder.text_frame
    tf.word_wrap = True
    tf.clear()
    if not bullets:
        return
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        run = p.add_run()
        run.text = text


# --------------------------------------------------------------------------
# Visual placeholders -> native chart / table / picture / icon row
# --------------------------------------------------------------------------


_MIN_CHART_HEIGHT_EMU = int(2.2 * 914400)   # ~2.2in — below this a chart just looks like a squashed sliver
_MIN_TABLE_ROW_HEIGHT_EMU = 320000          # ~0.35in per row, a legible minimum
_MIN_IMAGE_HEIGHT_EMU = int(2.3 * 914400)   # ~2.3in — same box-too-small problem as charts/tables


def _grow_box(box: tuple[int, int, int, int], design_system: DesignSystem, desired_height: int) -> tuple[int, int, int, int]:
    """Charts and tables need real vertical space. Several of this
    template's layouts ('1_Разделитель', '1_Визитка', ...) offer only a
    single-line subtitle placeholder for "body" content — fine for a short
    line of bullets, nowhere near tall enough for a 7-row table or a chart.
    python-pptx will happily create either at whatever tiny size you hand
    it; the shape then renders at its own real (much taller) natural size
    on the actual slide and overflows straight over whatever sits below it
    — which is exactly the "съезжает" overflow this fixes.

    Grows the placeholder's height toward ``desired_height``, extending
    downward within the slide's own bottom margin. Never shrinks a
    placeholder that was already big enough, and never pushes past the
    bottom margin even if that means falling short of ``desired_height``.
    """
    left, top, width, height = box
    bottom_margin = design_system.margins_emu.get("bottom", int(design_system.slide_height * 0.05))
    available_below = max(design_system.slide_height - top - bottom_margin, height)
    new_height = min(max(height, desired_height), available_below)
    return left, top, width, new_height


def _placeholder_box(placeholder) -> tuple[int, int, int, int]:
    return placeholder.left or 0, placeholder.top or 0, placeholder.width or 0, placeholder.height or 0


def _replace_with_chart(slide, placeholder, spec, accent_hex: str, design_system: DesignSystem) -> bool:
    left, top, width, height = _grow_box(_placeholder_box(placeholder), design_system, _MIN_CHART_HEIGHT_EMU)
    try:
        chart_data = CategoryChartData()
        chart_data.categories = spec.categories
        chart_data.add_series(spec.series_name, spec.values)
        xl_type = _CHART_TYPE_MAP.get(spec.chart_type, XL_CHART_TYPE.COLUMN_CLUSTERED)
        graphic_frame = slide.shapes.add_chart(xl_type, left, top, width, height, chart_data)
        chart = graphic_frame.chart
        chart.has_title = True
        chart.chart_title.text_frame.text = spec.series_name
        chart.has_legend = spec.chart_type == "pie"
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        plot = chart.plots[0]
        plot.has_data_labels = True
        try:
            for series in plot.series:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = RGBColor.from_string(accent_hex)
        except Exception:  # pie charts colour per-point, not per-series; not worth failing the slide over
            pass
    except Exception as exc:  # noqa: BLE001
        logger.warning("chart insertion failed, falling back to bullets: %s", exc)
        return False
    _remove_shape(placeholder)
    return True


def _replace_with_table(slide, placeholder, spec: TableSpec, accent_hex: str, design_system: DesignSystem) -> bool:
    rows, cols = len(spec.rows) + 1, len(spec.headers)
    desired_height = rows * _MIN_TABLE_ROW_HEIGHT_EMU
    left, top, width, height = _grow_box(_placeholder_box(placeholder), design_system, desired_height)
    try:
        shape = slide.shapes.add_table(rows, cols, left, top, width, height)
        table = shape.table
        for c, header in enumerate(spec.headers):
            cell = table.cell(0, c)
            cell.text = header
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(accent_hex)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor.from_string("FFFFFF")
        for r, row in enumerate(spec.rows, start=1):
            for c, value in enumerate(row):
                if c < cols:
                    table.cell(r, c).text = str(value)
    except Exception as exc:  # noqa: BLE001
        logger.warning("table insertion failed, falling back to bullets: %s", exc)
        return False
    _remove_shape(placeholder)
    return True


def _replace_with_picture(slide, box: tuple[int, int, int, int], img_bytes: bytes, placeholder=None) -> bool:
    """``box`` may come from a real placeholder OR from a stencil shape read
    straight off the layout (see ``_find_picture_stencil_box``) — the latter
    has no shape on the slide itself, so ``placeholder`` is None and there's
    nothing to remove afterward."""
    import io

    from PIL import Image

    left, top, width, height = box
    try:
        im = Image.open(io.BytesIO(img_bytes))
        iw, ih = im.size
        box_ratio = width / height if height else 1.0
        img_ratio = iw / ih if ih else 1.0
        if img_ratio > box_ratio:
            new_w, new_h = width, int(width / img_ratio)
        else:
            new_h, new_w = height, int(height * img_ratio)
        new_left = left + (width - new_w) // 2
        new_top = top + (height - new_h) // 2
        slide.shapes.add_picture(io.BytesIO(img_bytes), new_left, new_top, new_w, new_h)
    except Exception as exc:  # noqa: BLE001
        logger.warning("picture insertion failed: %s", exc)
        return False
    if placeholder is not None:
        _remove_shape(placeholder)
    return True


def _replace_with_icon_row(slide, placeholder, items: list[str], accent_hex: str) -> bool:
    try:
        add_icon_row(slide, placeholder.left, placeholder.top, placeholder.width, placeholder.height, items, accent_hex)
    except Exception as exc:  # noqa: BLE001
        logger.warning("icon row insertion failed, falling back to bullets: %s", exc)
        return False
    _remove_shape(placeholder)
    return True


# --------------------------------------------------------------------------
# Per-slide assembly
# --------------------------------------------------------------------------


def _fill_slide(
    slide,
    layout_spec: LayoutSpec,
    item: SlidePlanItem,
    *,
    llm: LLMClient,
    image_client: ImageClient,
    facts: list[SearchResult],
    language: str,
    accent_hex: str,
    design_system: DesignSystem,
    actual_layout,
) -> None:
    title_ph, bodies, picture_ph = _match_slide_placeholders(slide, layout_spec)
    if title_ph is not None:
        _set_title(title_ph, item.title or "")

    bullets_remaining = list(item.bullets)
    visual: VisualSpec = item.visual

    if visual.type == "chart" and bodies:
        target = bodies.pop(0)
        spec = build_chart_spec(llm, visual.hint, facts, visual.chart_type, language)
        if not (spec and _replace_with_chart(slide, target, spec, accent_hex, design_system)):
            bodies.insert(0, target)  # no usable data -> treat this box as a normal text well

    elif visual.type == "table" and bodies:
        target = bodies.pop(0)
        spec = build_table_spec(visual.hint, facts, language)
        if not (spec and _replace_with_table(slide, target, spec, accent_hex, design_system)):
            bodies.insert(0, target)

    elif visual.type == "image":
        box = None
        ph_to_remove = None
        if picture_ph is not None:
            box = _placeholder_box(picture_ph)
            ph_to_remove = picture_ph
        else:
            # No native PICTURE placeholder on this layout — check for the
            # template's "Вставить фото" stencil convention before falling
            # back to a text well (see _find_picture_stencil_box).
            stencil_box = _find_picture_stencil_box(actual_layout)
            if stencil_box is not None:
                box = stencil_box
            elif bodies:
                target = bodies[0]
                box = _placeholder_box(target)
                ph_to_remove = target

        if box is not None:
            # Same fix as charts/tables: a one-line subtitle box (or, less
            # often, an undersized stencil) is nowhere near tall enough for
            # a photo to read as a photo rather than a sliver.
            box = _grow_box(box, design_system, _MIN_IMAGE_HEIGHT_EMU)
            try:
                img_bytes = generate_image(image_client, visual.hint or item.title)
                if _replace_with_picture(slide, box, img_bytes, ph_to_remove):
                    if ph_to_remove is picture_ph:
                        picture_ph = None
                    elif ph_to_remove in bodies:
                        bodies.remove(ph_to_remove)
                # failure: nothing was removed, box's placeholder (if any) stays available below
            except Exception as exc:  # noqa: BLE001
                logger.warning("image generation failed: %s", exc)

    elif visual.type == "icon_row" and bodies:
        target = bodies.pop(0)
        if _replace_with_icon_row(slide, target, bullets_remaining or [item.title], accent_hex):
            bullets_remaining = []
        else:
            bodies.insert(0, target)

    if bodies:
        if len(bodies) >= 2 and bullets_remaining:
            mid = max(len(bullets_remaining) // 2, 1)
            _set_bullets(bodies[0], bullets_remaining[:mid])
            _set_bullets(bodies[1], bullets_remaining[mid:])
            for extra in bodies[2:]:
                _remove_shape(extra)
        elif bullets_remaining:
            _set_bullets(bodies[0], bullets_remaining)
            for extra in bodies[1:]:
                _remove_shape(extra)
        else:
            for extra in bodies:
                _remove_shape(extra)

    if picture_ph is not None and visual.type != "image":
        _remove_shape(picture_ph)

    if item.speaker_note:
        slide.notes_slide.notes_text_frame.text = item.speaker_note


# --------------------------------------------------------------------------
# Public entry point
# --------------------------------------------------------------------------


def build_variant_deck(
    template_path: str | Path,
    design_system: DesignSystem,
    plan: SlidePlan,
    facts: list[SearchResult],
    *,
    llm: LLMClient,
    image_client: ImageClient,
    language: str = "ru",
) -> Presentation:
    prs = Presentation(str(template_path))
    _strip_existing_slides(prs)
    accent_hex = _pick_accent(design_system)

    fallback_role_order = ("content_bullets", "generic", "two_content", "title_only", "blank")

    for item in plan.slides:
        layout_spec = design_system.best_layout_for(item.layout_role, fallback_roles=fallback_role_order)
        if layout_spec is None:
            layout_spec = design_system.layouts[0]
        actual_layout = prs.slide_masters[layout_spec.master_index].slide_layouts[layout_spec.layout_idx_in_master]
        slide = prs.slides.add_slide(actual_layout)
        _fill_slide(
            slide,
            layout_spec,
            item,
            llm=llm,
            image_client=image_client,
            facts=facts,
            language=language,
            accent_hex=accent_hex,
            design_system=design_system,
            actual_layout=actual_layout,
        )

    return prs
