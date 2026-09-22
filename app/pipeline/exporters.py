"""ТЗ §2 item 7: "Экспорт в .html, .pptx, .pdf."

.pptx is just ``Presentation.save`` — slide_builder already only ever
inserts native shapes (charts/tables as real objects, not flattened
images), so nothing extra is needed here for that requirement to hold.

.pdf goes through LibreOffice headless (the standard, high-fidelity way to
rasterize OOXML server-side; there is no pure-Python renderer that reliably
matches PowerPoint layout). It's an optional system dependency — if
``soffice`` isn't on PATH, ``export_pdf`` returns None and logs a warning
instead of failing the whole job.

.html is a from-scratch renderer: text stays real, selectable HTML
(positioned to match the .pptx geometry, EMU -> px at 96dpi), tables become
real ``<table>`` markup, and pictures embed as data URIs. Charts are the one
exception — reproducing OOXML chart rendering in a browser without shipping
a JS charting library would fight the "self-contained, no external scripts"
requirement, so chart *shapes* are redrawn with matplotlib from the same
data used to build the native pptx chart and embedded as a PNG. That's a
deliberate, documented trade-off (see README), not an oversight.
"""
from __future__ import annotations

import base64
import html
import logging
import subprocess
from pathlib import Path

from pptx import Presentation

from ..config import settings
from .template_parser import DesignSystem

logger = logging.getLogger(__name__)

EMU_PER_PX = 9525  # 96 dpi


def save_pptx(prs: Presentation, out_path: str | Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


def export_pdf(pptx_path: str | Path, out_dir: str | Path) -> Path | None:
    if not settings.pdf_export_enabled:
        return None
    pptx_path, out_dir = Path(pptx_path), Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    expected = out_dir / (pptx_path.stem + ".pdf")
    try:
        subprocess.run(
            [settings.soffice_binary, "--headless", "--norestore", "--convert-to", "pdf", "--outdir", str(out_dir), str(pptx_path)],
            check=True,
            capture_output=True,
            timeout=settings.pdf_export_timeout_seconds,
        )
    except FileNotFoundError:
        logger.warning("PDF export skipped: '%s' not found on PATH. Install LibreOffice to enable PDF export.", settings.soffice_binary)
        return None
    except subprocess.TimeoutExpired:
        logger.warning("PDF export timed out for %s", pptx_path)
        return None
    except subprocess.CalledProcessError as exc:
        logger.warning("PDF export failed for %s: %s", pptx_path, exc.stderr.decode(errors="replace")[:500])
        return None
    return expected if expected.exists() else None


# --------------------------------------------------------------------------
# HTML export
# --------------------------------------------------------------------------


def _emu_to_px(v: int | None) -> float:
    return round((v or 0) / EMU_PER_PX, 1)


def _run_style(font) -> str:
    parts = []
    try:
        if font.size is not None:
            parts.append(f"font-size:{font.size.pt}pt")
    except Exception:
        pass
    if font.bold:
        parts.append("font-weight:bold")
    if font.italic:
        parts.append("font-style:italic")
    try:
        if font.color and font.color.type is not None:
            parts.append(f"color:#{font.color.rgb}")
    except Exception:
        pass
    return ";".join(parts)


def _render_text_frame(tf) -> str:
    out = []
    for p in tf.paragraphs:
        text = "".join(html.escape(r.text) if not _run_style(r.font) else f'<span style="{_run_style(r.font)}">{html.escape(r.text)}</span>' for r in p.runs) if p.runs else "&nbsp;"
        indent = "list" if (p.level or 0) > 0 or len(tf.paragraphs) > 1 else "plain"
        bullet = "• " if indent == "list" and text.strip() and text != "&nbsp;" else ""
        align = {1: "center", 2: "right", 3: "justify"}.get(getattr(p.alignment, "value", None), "left")
        out.append(f'<div style="text-align:{align};margin:0 0 4px 0;">{bullet}{text}</div>')
    return "".join(out)


def _shape_fill_css(shape) -> str:
    try:
        fill = shape.fill
        if fill.type is not None and int(fill.type) == 1:
            return f"background-color:#{fill.fore_color.rgb};"
    except Exception:
        pass
    return ""


def _render_chart_png_b64(chart) -> str | None:
    try:
        import io

        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        plot = chart.plots[0]
        categories = [str(c) for c in plot.categories]
        series = list(plot.series)
        chart_type = chart.chart_type

        fig, ax = plt.subplots(figsize=(6, 3.3), dpi=110)
        if "PIE" in str(chart_type):
            values = list(series[0].values)
            ax.pie(values, labels=categories, autopct="%1.0f%%")
        elif "LINE" in str(chart_type):
            for s in series:
                ax.plot(categories, list(s.values), marker="o", label=s.name)
            ax.legend() if len(series) > 1 else None
        else:
            width = 0.8 / max(len(series), 1)
            for idx, s in enumerate(series):
                xs = [i + idx * width for i in range(len(categories))]
                ax.bar(xs, list(s.values), width=width, label=s.name)
            ax.set_xticks(range(len(categories)))
            ax.set_xticklabels(categories, rotation=15, ha="right")
            ax.legend() if len(series) > 1 else None
        fig.tight_layout()
        buf = io.BytesIO()
        fig.savefig(buf, format="png", transparent=True)
        plt.close(fig)
        return base64.b64encode(buf.getvalue()).decode("ascii")
    except Exception as exc:  # noqa: BLE001
        logger.warning("HTML export: could not render chart snapshot (%s)", exc)
        return None


def _render_shape(shape) -> str:
    left, top, width, height = _emu_to_px(shape.left), _emu_to_px(shape.top), _emu_to_px(shape.width), _emu_to_px(shape.height)
    base_style = f"position:absolute;left:{left}px;top:{top}px;width:{width}px;height:{height}px;"

    type_name = shape.shape_type.name if shape.shape_type is not None else ""

    if type_name == "PICTURE":
        try:
            image = shape.image
            b64 = base64.b64encode(image.blob).decode("ascii")
            return f'<img style="{base_style}object-fit:contain;" src="data:{image.content_type};base64,{b64}" alt="" />'
        except Exception:
            return ""

    if type_name == "CHART":
        b64 = _render_chart_png_b64(shape.chart)
        if b64:
            return f'<img style="{base_style}object-fit:contain;" src="data:image/png;base64,{b64}" alt="chart" />'
        return f'<div style="{base_style}display:flex;align-items:center;justify-content:center;color:#888;border:1px dashed #ccc;">[chart]</div>'

    if getattr(shape, "has_table", False) and shape.has_table:
        table = shape.table
        rows_html = []
        for r in table.rows:
            cells = "".join(f"<td style='border:1px solid #ddd;padding:4px 8px;'>{html.escape(c.text)}</td>" for c in r.cells)
            rows_html.append(f"<tr>{cells}</tr>")
        return f'<table style="{base_style}border-collapse:collapse;font-size:12pt;">{"".join(rows_html)}</table>'

    if type_name == "LINE":
        color = "#999999"
        try:
            color = f"#{shape.line.color.rgb}"
        except Exception:
            pass
        return f'<div style="{base_style}border-top:2px solid {color};"></div>'

    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        fill_css = _shape_fill_css(shape)
        text_color = ""
        if fill_css:
            text_color = "color:#fff;"
        inner = _render_text_frame(shape.text_frame)
        return f'<div style="{base_style}{fill_css}{text_color}padding:4px;box-sizing:border-box;overflow:hidden;">{inner}</div>'

    return ""


def _render_slide(slide, index: int, width_px: float, height_px: float, bg_hex: str) -> str:
    shapes_html = "".join(_render_shape(s) for s in slide.shapes)
    display = "block" if index == 0 else "none"
    return (
        f'<section class="slide" data-index="{index}" style="display:{display};position:relative;'
        f'width:{width_px}px;height:{height_px}px;background:#{bg_hex};margin:0 auto;overflow:hidden;">'
        f"{shapes_html}</section>"
    )


_HTML_SHELL = """<!DOCTYPE html>
<html lang="{lang}">
<head>
<meta charset="utf-8" />
<title>{title}</title>
<style>
  body {{ margin:0; padding:24px; background:#e9e9ef; font-family: '{font}', Arial, sans-serif; }}
  .deck {{ max-width: {width_px}px; margin: 0 auto; }}
  .slide {{ box-shadow: 0 2px 12px rgba(0,0,0,0.15); }}
  .nav {{ max-width: {width_px}px; margin: 12px auto; display:flex; gap:8px; align-items:center; font-family: Arial, sans-serif; }}
  .nav button {{ padding:6px 14px; cursor:pointer; }}
</style>
</head>
<body>
<div class="nav">
  <button onclick="go(-1)">&larr; Пред.</button>
  <span id="counter">1 / {n}</span>
  <button onclick="go(1)">След. &rarr;</button>
</div>
<div class="deck" id="deck">
{slides}
</div>
<script>
  let idx = 0;
  const slides = document.querySelectorAll('.slide');
  function go(delta) {{
    slides[idx].style.display = 'none';
    idx = Math.max(0, Math.min(slides.length - 1, idx + delta));
    slides[idx].style.display = 'block';
    document.getElementById('counter').textContent = (idx + 1) + ' / ' + slides.length;
  }}
  document.addEventListener('keydown', (e) => {{
    if (e.key === 'ArrowRight') go(1);
    if (e.key === 'ArrowLeft') go(-1);
  }});
</script>
</body>
</html>
"""


def export_html(prs: Presentation, design_system: DesignSystem, out_path: str | Path, title: str = "Presentation", language: str = "ru") -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    width_px = _emu_to_px(design_system.slide_width)
    height_px = _emu_to_px(design_system.slide_height)
    bg_hex = next((c.hex for c in design_system.colors if c.role == "lt1"), "FFFFFF")

    slides_html = "\n".join(_render_slide(s, i, width_px, height_px, bg_hex) for i, s in enumerate(prs.slides))
    doc = _HTML_SHELL.format(
        lang=language,
        title=html.escape(title),
        font=design_system.minor_font,
        width_px=int(width_px) + 4,
        n=len(prs.slides),
        slides=slides_html,
    )
    out_path.write_text(doc, encoding="utf-8")
    return out_path
