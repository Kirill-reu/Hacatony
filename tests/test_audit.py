from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu, Pt

from app.pipeline.audit import run_deterministic_audit
from app.pipeline.template_parser import parse_template


@pytest.fixture()
def design_system(tmp_path: Path):
    prs = Presentation()
    path = tmp_path / "template.pptx"
    prs.save(str(path))
    return parse_template(path)


def _save(prs: Presentation, tmp_path: Path, name: str) -> str:
    path = tmp_path / name
    prs.save(str(path))
    return str(path)


def test_flags_out_of_bounds_shape(tmp_path, design_system):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(-100000), Emu(0), Emu(500000), Emu(500000))
    box.text_frame.text = "off-slide"
    path = _save(prs, tmp_path, "d.pptx")
    report = run_deterministic_audit(path, design_system)
    assert any(i.code == "out_of_bounds" for i in report.issues)


def test_flags_overlap(tmp_path, design_system):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    b1 = slide.shapes.add_textbox(Emu(500000), Emu(500000), Emu(1000000), Emu(1000000))
    b1.text_frame.text = "one"
    b2 = slide.shapes.add_textbox(Emu(700000), Emu(700000), Emu(1000000), Emu(1000000))
    b2.text_frame.text = "two"
    path = _save(prs, tmp_path, "d.pptx")
    report = run_deterministic_audit(path, design_system)
    assert any(i.code == "overlap" for i in report.issues)


def test_flags_too_many_bullets(tmp_path, design_system):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(500000), Emu(500000), Emu(4000000), Emu(3000000))
    tf = box.text_frame
    for i in range(8):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.add_run().text = f"bullet {i}"
    path = _save(prs, tmp_path, "d.pptx")
    report = run_deterministic_audit(path, design_system)
    assert any(i.code == "too_many_bullets" for i in report.issues)


def test_flags_leftover_placeholder_text(tmp_path, design_system):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(500000), Emu(500000), Emu(2000000), Emu(500000))
    box.text_frame.text = "TODO: заполнить позже"
    path = _save(prs, tmp_path, "d.pptx")
    report = run_deterministic_audit(path, design_system)
    assert any(i.code == "leftover_placeholder" for i in report.issues)


def test_flags_empty_slide(tmp_path, design_system):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])  # blank, no shapes at all
    path = _save(prs, tmp_path, "d.pptx")
    report = run_deterministic_audit(path, design_system)
    assert any(i.code == "empty_slide" for i in report.issues)


def test_flags_duplicate_slides(tmp_path, design_system):
    prs = Presentation()
    for _ in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Emu(500000), Emu(500000), Emu(2000000), Emu(500000))
        box.text_frame.text = "Одинаковый слайд"
    path = _save(prs, tmp_path, "d.pptx")
    report = run_deterministic_audit(path, design_system)
    assert any(i.code == "duplicate_slide" for i in report.issues)


def test_clean_slide_has_no_issues(tmp_path, design_system):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.placeholders[0].text_frame.text = "Нормальный заголовок"
    tf = slide.placeholders[1].text_frame
    tf.clear()
    tf.paragraphs[0].add_run().text = "Один аккуратный буллет"
    path = _save(prs, tmp_path, "d.pptx")
    report = run_deterministic_audit(path, design_system)
    assert report.issues == []
