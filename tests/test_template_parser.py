"""Builds a throwaway .pptx (python-pptx's own default template, which ships
the standard set of layouts: Title Slide, Title and Content, Section Header,
Two Content, Comparison, Title Only, Blank, Content with Caption, Picture
with Caption, ...) and checks template_parser decomposes it sensibly. No
network, no fixtures on disk needed."""
from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from app.pipeline.template_parser import parse_template


@pytest.fixture()
def sample_template(tmp_path: Path) -> Path:
    prs = Presentation()
    for i in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.placeholders[0].text_frame.text = f"Sample slide {i}"
        slide.placeholders[1].text_frame.text = "Some body text for typography extraction"
    path = tmp_path / "template.pptx"
    prs.save(str(path))
    return path


def test_parses_theme_colors_and_fonts(sample_template: Path):
    ds = parse_template(sample_template)
    assert ds.major_font
    assert ds.minor_font
    roles = {c.role for c in ds.colors}
    assert {"dk1", "lt1", "accent1"}.issubset(roles)
    for c in ds.colors:
        assert len(c.hex) == 6
        int(c.hex, 16)  # valid hex


def test_classifies_standard_layouts(sample_template: Path):
    ds = parse_template(sample_template)
    assert ds.layout_roles.get("title_slide")
    assert ds.layout_roles.get("content_bullets")
    assert ds.layout_roles.get("two_content")
    assert ds.layout_roles.get("content_image")
    assert ds.layout_roles.get("section_header")


def test_type_scale_is_nonempty_and_sorted(sample_template: Path):
    ds = parse_template(sample_template)
    assert ds.type_scale_pt == sorted(ds.type_scale_pt)
    assert len(ds.type_scale_pt) > 0


def test_best_layout_for_falls_back(sample_template: Path):
    ds = parse_template(sample_template)
    # "generic" almost never exists on the stock template; fallback must still resolve to something.
    layout = ds.best_layout_for("generic", fallback_roles=("content_bullets", "blank"))
    assert layout is not None


def test_margins_are_within_slide_bounds(sample_template: Path):
    ds = parse_template(sample_template)
    assert 0 <= ds.margins_emu["left"] < ds.slide_width
    assert 0 <= ds.margins_emu["top"] < ds.slide_height
